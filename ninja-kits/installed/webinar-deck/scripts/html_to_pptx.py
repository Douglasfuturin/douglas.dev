#!/usr/bin/env python3
"""
html_to_pptx.py, convert a rendered deck.html into a .pptx.

Each slide is screenshotted at native 16:9 via headless Chromium (Playwright)
and placed full-bleed onto a 16:9 PowerPoint slide. The result is pixel-perfect
and presents in PowerPoint or Keynote. Text is NOT editable in the PPTX, to
change content, edit the config.json and re-render, then re-export.

Works for a deck of any slide count, it reads the actual number of slides
from the rendered deck.

Usage:
  python3 html_to_pptx.py --deck workspace/my-deck/deck.html
  python3 html_to_pptx.py --deck .../deck.html --output .../deck.pptx

Requires: playwright (+ chromium browser), python-pptx.
  pip install -r requirements.txt
  playwright install chromium
"""

import argparse
import shutil
import sys
import tempfile
from pathlib import Path


def main():
    ap = argparse.ArgumentParser(description="Convert a rendered deck.html into a .pptx")
    ap.add_argument("--deck", required=True, help="Path to the rendered deck.html")
    ap.add_argument("--output", help="Output .pptx path (default: deck.pptx beside the deck)")
    ap.add_argument("--width", type=int, default=1920, help="Capture width (default 1920)")
    ap.add_argument("--height", type=int, default=1080, help="Capture height (default 1080)")
    args = ap.parse_args()

    deck = Path(args.deck).resolve()
    if not deck.exists():
        print(f"ERROR: deck not found: {deck}", file=sys.stderr)
        sys.exit(1)
    out = Path(args.output).resolve() if args.output else deck.with_suffix(".pptx")

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("ERROR: playwright not installed.\n"
              "  pip install playwright python-pptx\n"
              "  playwright install chromium", file=sys.stderr)
        sys.exit(1)
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        print("ERROR: python-pptx not installed.\n  pip install python-pptx", file=sys.stderr)
        sys.exit(1)

    shots_dir = Path(tempfile.mkdtemp(prefix="deck-pptx-"))
    shot_paths = []
    try:
        with sync_playwright() as pw:
            try:
                browser = pw.chromium.launch()
            except Exception as e:  # noqa: BLE001
                print(f"ERROR: could not launch Chromium, run `playwright install chromium`.\n  {e}",
                      file=sys.stderr)
                sys.exit(1)
            page = browser.new_page(viewport={"width": args.width, "height": args.height})
            page.goto(deck.as_uri())
            page.wait_for_load_state("networkidle")

            total = page.evaluate("document.querySelectorAll('.slide').length")
            if not total:
                print("ERROR: no .slide elements found in the deck.", file=sys.stderr)
                sys.exit(1)
            has_hook = page.evaluate("typeof window.__gotoSlide === 'function'")
            if not has_hook:
                print("ERROR: deck has no __gotoSlide hook, re-render with the current "
                      "deck-template.html.", file=sys.stderr)
                sys.exit(1)

            print(f"Capturing {total} slides at {args.width}x{args.height}...", file=sys.stderr)
            canvas = page.locator(".slide-canvas")
            for i in range(total):
                page.evaluate(f"window.__gotoSlide({i})")
                page.wait_for_timeout(350)  # the slide fade is 200ms, let it settle
                shot = shots_dir / f"slide-{i + 1:04d}.png"
                canvas.screenshot(path=str(shot))
                shot_paths.append(shot)
                if (i + 1) % 25 == 0 or i + 1 == total:
                    print(f"  {i + 1}/{total}", file=sys.stderr)
            browser.close()

        # Assemble the .pptx, one full-bleed picture per 16:9 slide
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        for shot in shot_paths:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(shot), 0, 0,
                                     width=prs.slide_width, height=prs.slide_height)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        print(f"Done: {out}  ({len(shot_paths)} slides)", file=sys.stderr)
    finally:
        shutil.rmtree(shots_dir, ignore_errors=True)


if __name__ == "__main__":
    main()
